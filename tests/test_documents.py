import os

import pytest

import engine.documents as documents
from engine.documents import (
    SUPPORTED_SUFFIXES,
    UnsupportedDocumentError,
    convert_to_pdf,
    text_blocks,
)


def _pdf_text(path):
    import pymupdf

    doc = pymupdf.open(path)
    try:
        return "".join(page.get_text() for page in doc)
    finally:
        doc.close()


def test_convert_txt_to_pdf_next_to_source(tmp_path):
    source = tmp_path / "notes.txt"
    source.write_text("Alex Rivera\nParalegal\n", encoding="utf-8")

    out = convert_to_pdf(str(source))

    assert out == os.path.realpath(str(tmp_path / "notes.pdf"))
    assert os.path.exists(out)
    assert "Alex Rivera" in _pdf_text(out)


def test_convert_markdown_to_pdf_styles_headings_and_bullets(tmp_path):
    source = tmp_path / "resume.md"
    source.write_text(
        "# Alex Rivera\n\n- Westlaw\n- e-discovery\n\nLitigation support.\n",
        encoding="utf-8",
    )

    out = convert_to_pdf(str(source), str(tmp_path / "out" / "resume.pdf"))

    assert out == os.path.realpath(str(tmp_path / "out" / "resume.pdf"))
    text = _pdf_text(out)
    assert "Alex Rivera" in text
    assert "Westlaw" in text
    assert "Litigation support." in text


def test_convert_pdf_passthrough_returns_source(tmp_path):
    source = tmp_path / "report.pdf"
    source.write_bytes(b"%PDF-1.4 placeholder")

    assert convert_to_pdf(str(source)) == os.path.realpath(str(source))


def test_convert_pdf_copies_to_explicit_target(tmp_path):
    source = tmp_path / "report.pdf"
    source.write_bytes(b"%PDF-1.4 placeholder")
    target = tmp_path / "copy" / "report.pdf"

    out = convert_to_pdf(str(source), str(target))

    assert out == os.path.realpath(str(target))
    assert target.read_bytes() == source.read_bytes()


def test_convert_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        convert_to_pdf(str(tmp_path / "missing.txt"))


def test_convert_unsupported_suffix_raises(tmp_path):
    source = tmp_path / "image.png"
    source.write_bytes(b"\x89PNG")

    with pytest.raises(UnsupportedDocumentError) as error:
        convert_to_pdf(str(source))

    assert "Unsupported file type" in str(error.value)
    assert ".docx" in str(error.value)


def test_convert_doc_reports_word_requirement(monkeypatch, tmp_path):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"\xd0\xcf\x11\xe0 legacy word")

    def fail(source_path, pdf_path):
        raise ImportError("docx2pdf is not installed")

    monkeypatch.setattr(documents, "convert_with_word_bridge", fail)

    with pytest.raises(RuntimeError) as error:
        convert_to_pdf(str(source))

    assert "docx2pdf" in str(error.value)


def test_convert_docx_delegates_to_docx_converter(tmp_path):
    import docx

    document = docx.Document()
    document.add_heading("Resume", level=1)
    document.add_paragraph("Alex Rivera")
    source = tmp_path / "resume.docx"
    document.save(str(source))

    out = convert_to_pdf(str(source))

    assert out == os.path.realpath(str(tmp_path / "resume.pdf"))
    assert "Alex Rivera" in _pdf_text(out)


def test_text_blocks_splits_headings_bullets_and_body():
    blocks = text_blocks("# Title\n\n- One\n* Two\nPlain line\n\n")

    assert blocks == [
        ("heading", "Title"),
        ("bullet", "One"),
        ("bullet", "Two"),
        ("body", "Plain line"),
    ]


def test_supported_suffixes_cover_documents_and_text():
    for suffix in (".docx", ".doc", ".txt", ".md", ".pdf"):
        assert suffix in SUPPORTED_SUFFIXES
