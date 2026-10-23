import os

import pytest

from engine.docx import convert_docx_to_pdf, safe_pdf_name


def _make_docx(path, paragraphs=None):
    import docx

    document = docx.Document()
    document.add_heading("Test Resume", level=1)
    for text in paragraphs or ["First paragraph", "Second paragraph"]:
        document.add_paragraph(text)
    document.save(path)
    return path


def test_convert_writes_pdf_next_to_docx(tmp_path):
    docx_path = _make_docx(str(tmp_path / "resume.docx"))
    out = convert_docx_to_pdf(docx_path)

    assert out == str(tmp_path / "resume.pdf")
    assert os.path.exists(out)


def test_convert_to_custom_pdf_path(tmp_path):
    docx_path = _make_docx(str(tmp_path / "resume.docx"))
    target = str(tmp_path / "out" / "custom.pdf")
    out = convert_docx_to_pdf(docx_path, target)

    assert out == target
    assert os.path.exists(out)


def test_convert_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        convert_docx_to_pdf(str(tmp_path / "nope.docx"))


def test_convert_wrong_extension_raises(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_text("not a docx")
    with pytest.raises(ValueError):
        convert_docx_to_pdf(str(path))


def test_convert_preserves_table_text(tmp_path):
    import docx
    import pymupdf

    docx_path = str(tmp_path / "with_table.docx")
    document = docx.Document()
    document.add_heading("Resume", level=1)
    document.add_paragraph("A short paragraph about me.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python React SQL"
    document.save(docx_path)

    out = convert_docx_to_pdf(docx_path)
    doc = pymupdf.open(out)
    try:
        text = "".join(page.get_text() for page in doc)
    finally:
        doc.close()

    assert "Python React SQL" in text
    assert "Skills" in text


def test_safe_pdf_name_defaults_next_to_docx(tmp_path):
    docx_path = str(tmp_path / "resume.docx")
    assert safe_pdf_name(docx_path) == os.path.realpath(str(tmp_path / "resume.pdf"))
