import pymupdf

from engine.scorer import (
    _MAX_JOB_KEYWORDS,
    _extract_skills_section,
    _flow_pdf_block,
    _job_cover,
    _job_priority_keywords,
    _read_resume_text,
    determine_badge,
    extract_job_keywords,
    extract_resume_keywords,
    normalize_synonyms,
    score_jobs,
    score_jobs_by_track,
)


def _job(title, desc, url="http://example.com/job"):
    return {"title": title, "company": "Acme", "location": "NYC", "description": desc, "url": url}


def test_zero_keyword_job_scores_zero():
    jobs = [_job("Generic", "generic responsibilities and duties")]
    scored, _ = score_jobs(jobs, "", ["python", "react", "sql"], 1.0)
    assert scored[0]["match_score"] == 0


def test_monotonicity_more_matches_scores_higher():
    skills = ["python", "react", "sql"]
    low = _job("Low", "react experience")
    high = _job("High", "python react sql experience")
    scored, _ = score_jobs([low, high], "", skills, 1.0)
    by_title = {j["title"]: j["match_score"] for j in scored}
    assert by_title["High"] > by_title["Low"]
    assert by_title["Low"] > 0


def test_non_tech_domain_fixture():
    skills = ["patient care", "HIPAA", "medication administration"]
    jobs = [
        _job("RN", "provide patient care and administer medication"),
        _job("CNA", "assist with daily living"),
    ]
    scored, _ = score_jobs(jobs, "", skills, 1.0)
    by_title = {j["title"]: j["match_score"] for j in scored}
    assert by_title["RN"] >= by_title["CNA"]


def test_determine_badge():
    assert determine_badge("requires a cover letter") == "Amber"
    assert determine_badge("standard application") == "Slate"
    assert determine_badge("") == "Slate"


def _make_resume_pdf(path, text="Python React SQL Westlaw"):
    doc = pymupdf.open()
    page = doc.new_page()
    for index, line in enumerate(text.splitlines() or [text]):
        page.insert_text((72, 72 + index * 14), line)
    doc.save(str(path))
    doc.close()


def _make_resume_docx(path, paragraphs):
    import docx

    document = docx.Document()
    document.add_heading("Test Resume", level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))
    return str(path)


def test_extract_resume_keywords_from_pdf(tmp_path):
    pdf = tmp_path / "resume.pdf"
    _make_resume_pdf(pdf)
    kw = extract_resume_keywords(str(pdf), [])
    assert "Python" in kw
    assert "SQL" in kw


def test_extract_keywords_reads_docx(tmp_path):
    docx_path = _make_resume_docx(
        tmp_path / "resume.docx",
        [
            "Skills",
            "Python, SQL, Docker, Figma, Information Architecture",
            "Wrote API documentation and style guides for a developer platform",
        ],
    )
    kw = extract_resume_keywords(str(docx_path), [])
    lowered = [k.lower() for k in kw]

    for expected in ("python", "sql", "docker", "figma", "information architecture"):
        assert expected in lowered
    assert all(k.count("(") == k.count(")") for k in kw)


def test_pdf_and_docx_produce_equivalent_keywords(tmp_path):
    from engine.docx import convert_docx_to_pdf

    docx_path = _make_resume_docx(
        tmp_path / "resume.docx",
        [
            "Skills",
            "Python, SQL, Docker, Figma, Behavior Management",
            "Research (Concepts, IA), Information Architecture",
            "Documentation, technical writing, content design, API reference",
        ],
    )
    pdf_path = convert_docx_to_pdf(docx_path)

    from_pdf = extract_resume_keywords(pdf_path, ["Python", "SQL"])
    from_docx = extract_resume_keywords(docx_path, ["Python", "SQL"])

    pdf_set = {k.lower() for k in from_pdf}
    docx_set = {k.lower() for k in from_docx}
    sym_diff = pdf_set ^ docx_set
    union = pdf_set | docx_set

    assert len(sym_diff) <= max(2, 0.15 * len(union))
    for must_have in ("Python", "SQL"):
        assert must_have in from_pdf
        assert must_have in from_docx


def test_pdf_visual_wrap_keeps_multiword_skill_and_separate_bullets(tmp_path):
    import docx

    docx_path = tmp_path / "resume.docx"
    document = docx.Document()
    document.add_paragraph("Skills")
    document.add_paragraph("Information Architecture, Accessibility")
    document.save(docx_path)

    pdf_path = tmp_path / "resume.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Skills", fontsize=11)
    page.insert_textbox(
        pymupdf.Rect(72, 88, 156, 150),
        "Information Architecture, Accessibility",
        fontsize=10,
    )
    doc.save(pdf_path)
    doc.close()

    pdf_terms = _extract_skills_section(_read_resume_text(str(pdf_path)).lower())
    docx_terms = _extract_skills_section(_read_resume_text(str(docx_path)).lower())

    assert pdf_terms == docx_terms == ["information architecture", "accessibility"]


def test_pdf_block_flow_does_not_merge_separate_bullets():
    def line(text, y):
        return {
            "bbox": (72, y, 180, y + 10),
            "spans": [{"text": text, "size": 10}],
        }

    flowed = _flow_pdf_block([
        line("• Information", 10),
        line("Architecture", 21),
        line("• Accessibility", 32),
    ])

    assert flowed == "• Information Architecture\n• Accessibility"


def test_skills_section_keeps_parenthesised_terms(tmp_path):
    pdf = tmp_path / "resume.pdf"
    _make_resume_pdf(
        pdf,
        "Skills\nResearch (Concepts, IA), Figma, Docker\nWorked on design systems",
    )
    kw = extract_resume_keywords(str(pdf), [])
    lowered = [k.lower() for k in kw]

    assert "figma" in lowered
    assert "docker" in lowered
    assert "(concepts)" not in lowered
    assert "(indesign" not in lowered
    assert not any(k.endswith(")") for k in kw)


def test_junk_terms_are_dropped(tmp_path):
    pdf = tmp_path / "resume.pdf"
    _make_resume_pdf(
        pdf,
        "Skills\nAnd child abuse prevention, Additional skills\nIssued August, "
        "April 2024 - Present, September, January, Remote, Freelance\n"
        "Languages: fluent in Spanish",
    )
    lowered = [k.lower() for k in extract_resume_keywords(str(pdf), [])]

    for junk in (
        "august",
        "september",
        "january",
        "issued august",
        "april 2024 - present",
        "additional skills",
        "and child abuse prevention",
        "remote",
        "freelance",
        "languages: fluent in spanish",
    ):
        assert junk not in lowered


def test_no_case_duplicate_keywords(tmp_path):
    pdf = tmp_path / "resume.pdf"
    _make_resume_pdf(
        pdf,
        "Skills\nCSS, css, Docker, docker, GIT, git, Javascript, JavaScript, API, api",
    )
    lowered = [k.lower() for k in extract_resume_keywords(str(pdf), [])]

    assert len(lowered) == len(set(lowered))


def _tracks(tech_pdf, ux_pdf):
    return [
        {"name": "Tech Writer", "color": "#1D4ED8", "route_keywords": ["python", "api"], "path": str(tech_pdf)},
        {"name": "UX Writer", "color": "#7E22CE", "route_keywords": ["ux", "figma"], "path": str(ux_pdf)},
    ]


def test_score_jobs_by_track_different_resumes_score_differently(tmp_path):
    tech_pdf = tmp_path / "tech.pdf"
    ux_pdf = tmp_path / "ux.pdf"
    _make_resume_pdf(tech_pdf, "Python React SQL")
    _make_resume_pdf(ux_pdf, "Figma UX Research")

    jobs = [
        _job("Tech", "python api documentation"),
        _job("UX", "ux research"),
    ]
    scored, keywords = score_jobs_by_track(jobs, _tracks(tech_pdf, ux_pdf), str(tech_pdf), [], 1.0)

    by_title = {j["title"]: j for j in scored}
    assert by_title["Tech"]["resume_track"] == "Tech Writer"
    assert by_title["UX"]["resume_track"] == "UX Writer"
    assert by_title["Tech"]["match_score"] > 0
    assert by_title["Tech"]["match_score"] != by_title["UX"]["match_score"]
    # Both resumes' keywords appear in the union.
    assert "Python" in keywords
    assert "Figma" in keywords


def test_score_jobs_by_track_track_without_path_falls_back_to_default(tmp_path):
    default_pdf = tmp_path / "default.pdf"
    _make_resume_pdf(default_pdf, "Python React SQL")
    tracks = [{"name": "Only Track", "color": "#111111", "route_keywords": ["python"]}]

    scored, keywords = score_jobs_by_track(
        [_job("Dev", "python developer")], tracks, str(default_pdf), [], 1.0
    )
    assert scored[0]["resume_track"] == "Only Track"
    assert scored[0]["resume_color"] == "#111111"
    assert scored[0]["match_score"] > 0
    assert "Python" in keywords


def test_score_jobs_by_track_routes_by_route_keywords(tmp_path):
    a_pdf = tmp_path / "a.pdf"
    b_pdf = tmp_path / "b.pdf"
    _make_resume_pdf(a_pdf, "Python")
    _make_resume_pdf(b_pdf, "Figma")
    tracks = [
        {"name": "A", "color": "#AAAAAA", "route_keywords": ["python"], "path": str(a_pdf)},
        {"name": "B", "color": "#BBBBBB", "route_keywords": ["figma"], "path": str(b_pdf)},
    ]

    scored, _ = score_jobs_by_track([_job("Designer", "figma prototyping")], tracks, str(a_pdf), [], 1.0)
    assert scored[0]["resume_track"] == "B"
    assert scored[0]["resume_color"] == "#BBBBBB"


def test_job_coverage_rewards_requirement_coverage(tmp_path):
    resume_pdf = tmp_path / "resume.pdf"
    _make_resume_pdf(resume_pdf, "Python React SQL")

    covered = _job("Covered", "python react sql")
    partly = _job(
        "Partly", "python react sql kubernetes terraform ansible jenkins"
    )
    scored, keywords = score_jobs([covered, partly], str(resume_pdf), [], 1.0)
    by_title = {j["title"]: j for j in scored}

    # Both postings repeat every résumé keyword, so the résumé side ties.
    assert len(by_title["Covered"]["key_matches"]) == 3
    assert len(by_title["Partly"]["key_matches"]) == 3
    assert by_title["Covered"]["job_coverage"] > by_title["Partly"]["job_coverage"]
    assert by_title["Covered"]["match_score"] > by_title["Partly"]["match_score"]
    assert 0 <= by_title["Partly"]["job_coverage"] <= 100
    assert len(by_title["Partly"]["job_keywords"]) <= _MAX_JOB_KEYWORDS


def test_job_coverage_matches_regular_inflections_and_common_aliases():
    coverage = _job_cover(
        ["assessments", "API", "UX"],
        "Assessment planning, application programming interfaces, "
        "and user experience research.",
    )

    assert coverage == 1.0


def test_normalize_synonyms_accepts_groups_and_mappings():
    from_mapping = normalize_synonyms({"docs-as-code": ["documentation as code"]})
    from_groups = normalize_synonyms([["docs-as-code", "documentation as code"]])

    assert from_mapping == from_groups
    assert from_groups["docs-as-code"] == ("documentation as code",)
    assert from_groups["documentation as code"] == ("docs-as-code",)
    # A lone term carries no equivalence, and blanks never reach the map.
    assert normalize_synonyms([["only-term"]]) == {}
    assert normalize_synonyms([["  ", "term"]]) == {}


def test_profile_synonyms_match_in_both_directions(tmp_path):
    resume_pdf = tmp_path / "resume.pdf"
    _make_resume_pdf(resume_pdf, "Skills\ndocs-as-code, contract review")
    synonyms = normalize_synonyms([["docs-as-code", "documentation as code"]])

    matching = _job("Docs", "Requirements\ndocumentation as code")
    unrelated = _job("Legal", "Requirements\nparalegal support")

    scored, _ = score_jobs([matching, unrelated], str(resume_pdf), [], 1.0, synonyms=synonyms)
    by_title = {j["title"]: j for j in scored}

    # The résumé says "docs-as-code" and the posting says "documentation as code".
    assert "docs-as-code" in by_title["Docs"]["key_matches"]
    assert by_title["Docs"]["job_coverage"] > by_title["Legal"]["job_coverage"]
    assert by_title["Docs"]["match_score"] > by_title["Legal"]["match_score"]


def test_synonyms_do_not_loosen_whole_phrase_matching(tmp_path):
    resume_pdf = tmp_path / "resume.pdf"
    _make_resume_pdf(resume_pdf, "Skills\ncontent design")
    synonyms = normalize_synonyms([["content design", "ux writing"]])

    partial = _job("Partial", "Requirements\nux writer")
    exact = _job("Exact", "Requirements\nux writing")

    scored, _ = score_jobs([partial, exact], str(resume_pdf), [], 1.0, synonyms=synonyms)
    by_title = {j["title"]: j for j in scored}

    assert by_title["Exact"]["job_coverage"] > by_title["Partial"]["job_coverage"]


def test_required_keywords_outweigh_preferred_keywords():
    job = _job(
        "Technical Writer",
        "Requirements\nPython, SQL\n"
        "Preferred Qualifications\nFigma, UX research\n"
        "Responsibilities\nWrite product documentation",
    )
    required, preferred = _job_priority_keywords(job)
    keywords = extract_job_keywords(job)

    assert {term.lower() for term in required} >= {"python", "sql"}
    assert {term.lower() for term in preferred} >= {"figma", "ux research"}
    assert _job_cover(keywords, "python", required, preferred) > _job_cover(
        keywords, "figma", required, preferred
    )


def test_posting_without_text_falls_back_to_resume_overlap():
    # "go" is too short for the posting extractor, so the posting side yields
    # nothing and the score stays the résumé-side overlap it was before.
    scored, _ = score_jobs([_job("", "go")], "", ["Go", "python"], 1.0)
    result = scored[0]

    assert result["job_keywords"] == []
    assert result["job_coverage"] == 0
    assert result["match_score"] == int(1 / 2 * 100)


def test_title_only_posting_cannot_claim_full_requirement_coverage(tmp_path):
    # Boards such as Built In return postings without a body, so the only
    # extractable "requirements" are the words of the title. Full coverage from a
    # title alone would hand every such posting the same top score.
    resume_pdf = tmp_path / "resume.pdf"
    _make_resume_pdf(
        resume_pdf,
        "Skills\nPython, SQL, React, Docker, Kubernetes, Figma, Jira, Confluence,"
        " Markdown, Git, Postman, Swagger\nExperience\nTechnical Writer",
    )

    scored, _ = score_jobs(
        [_job("Technical Writer", "")], str(resume_pdf), ["technical writer"], 1.5
    )
    result = scored[0]

    assert result["job_keywords"]
    assert result["job_coverage"] == 100
    assert result["match_score"] < 50


def test_score_is_clamped_and_bounded(tmp_path):
    resume_pdf = tmp_path / "resume.pdf"
    _make_resume_pdf(resume_pdf, "Python React SQL")

    jobs = [_job("Perfect", "python react sql"), _job("None", "nothing here at all")]
    scored, _ = score_jobs(jobs, str(resume_pdf), [], 5.0)
    scores = [j["match_score"] for j in scored]

    assert all(0 <= score <= 100 for score in scores)
    assert max(scores) == 100


def test_job_keywords_capped():
    description = " ".join(f"skillterm{index}" for index in range(40))
    keywords = extract_job_keywords(_job("Engineer", description))

    assert len(keywords) == _MAX_JOB_KEYWORDS


def test_one_resume_read_per_sweep(tmp_path, monkeypatch):
    import engine.scorer as scorer

    tech_pdf = tmp_path / "tech.pdf"
    ux_pdf = tmp_path / "ux.pdf"
    _make_resume_pdf(tech_pdf, "Python React SQL")
    _make_resume_pdf(ux_pdf, "Figma UX Research")

    reads = []
    real_read = scorer._read_resume_text

    def counting_read(path):
        reads.append(path)
        return real_read(path)

    monkeypatch.setattr(scorer, "_read_resume_text", counting_read)

    jobs = [_job(f"Job {index}", "python react sql figma research") for index in range(5)]
    scorer.score_jobs([dict(job) for job in jobs], str(tech_pdf), [], 1.0)
    assert len(reads) == 1

    reads.clear()
    scorer.score_jobs_by_track([dict(job) for job in jobs], _tracks(tech_pdf, ux_pdf), str(tech_pdf), [], 1.0)
    # One read per unique track path, not one per posting.
    assert len(reads) == 2
